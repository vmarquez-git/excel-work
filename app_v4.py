import re
import logging
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook


RE_YM = re.compile(r"^(?P<ym>\d{4}-\d{2})\b")
RE_YM_DASH = re.compile(r"^\d{4}-\d{2}$")

SHEET_OB = "Snapshot OB TimePhased"
SHEET_CB = "Snapshot CB TimePhased"
SHEET_CHANGE = "Snapshot Change TimePhased"
SHEET_CHANGE_DETAILS = "Snapshot Change Details"
SHEET_ACTUALS = "Snapshot Actuals"
SHEET_ACTUALS_ETC = "Snapshot Actuals plus ETC"

REQUIRED_SHEETS = {
    SHEET_OB,
    SHEET_CB,
    SHEET_CHANGE,
    SHEET_CHANGE_DETAILS,
    SHEET_ACTUALS,
    SHEET_ACTUALS_ETC,
}


def setup_logging(out_dir: Path) -> Path:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_path = out_dir / "run.log"

    logger = logging.getLogger()
    logger.setLevel(logging.INFO)
    logger.handlers.clear()

    fmt = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")

    fh = logging.FileHandler(log_path, encoding="utf-8")
    fh.setFormatter(fmt)
    logger.addHandler(fh)

    sh = logging.StreamHandler()
    sh.setFormatter(fmt)
    logger.addHandler(sh)

    return log_path


def validate_required_sheets(file_path: Path) -> None:
    """
    Ensure the workbook contains the required sheets.
    """
    file_path = Path(file_path)
    logging.info("Checking required sheets in %s", file_path.name)

    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        found_sheets = set(wb.sheetnames)
        missing = REQUIRED_SHEETS - found_sheets

        if missing:
            raise RuntimeError(
                f"{file_path.name} is missing required sheets: {', '.join(sorted(missing))}. "
                f"Found sheets: {', '.join(sorted(found_sheets))}"
            )

        logging.info("Required sheets validated in %s", file_path.name)
    finally:
        wb.close()


def _to_month(ym: str) -> pd.Timestamp:
    if not RE_YM_DASH.match(ym):
        raise ValueError(f"Invalid month '{ym}'. Expected YYYY-MM.")
    return pd.to_datetime(ym + "-01")


def _clean_cols(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = (
        df.columns.astype(str)
        .str.replace("\u00A0", " ", regex=False)
        .str.strip()
    )
    ren = {
        "Project name": "Project Name",
        "ProjectName": "Project Name",
        "Project id": "Project ID",
        "Project Id": "Project ID",
        "ProjectID": "Project ID",
    }
    df = df.rename(columns={k: v for k, v in ren.items() if k in df.columns})
    return df


def _read_sheet(file_path: Path, sheet_name: str) -> pd.DataFrame:
    """
    Read a worksheet by name using pandas. Row 1 is assumed to be the header.
    """
    file_path = Path(file_path)

    df = pd.read_excel(
        file_path,
        sheet_name=sheet_name,
        header=0,
        engine="openpyxl",
    )

    if df.empty:
        logging.warning("Sheet '%s' in %s is empty or has no data rows", sheet_name, file_path.name)

    return _clean_cols(df)


def _month_cols(df: pd.DataFrame, suffix: str | tuple[str, ...]) -> list[str]:
    cols: list[str] = []
    for c in df.columns:
        s = str(c)
        if RE_YM.match(s) and s.endswith(suffix):
            cols.append(s)
    return cols


def _sum_month_cols_by_project(df: pd.DataFrame, cols: Iterable[str], value_name: str) -> pd.DataFrame:
    df = df.copy()
    need = {"Project ID", "Project Name"}
    if not need.issubset(df.columns):
        raise ValueError(f"Missing keys {need}. Found: {list(df.columns)}")

    cols = [c for c in cols if c in df.columns]
    if not cols:
        logging.warning("No month columns available for %s", value_name)
        return pd.DataFrame(columns=["Project ID", "Project Name", value_name])

    for c in cols:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)

    out = df.groupby(["Project ID", "Project Name"], as_index=False)[cols].sum()
    out[value_name] = out[cols].sum(axis=1)
    return out[["Project ID", "Project Name", value_name]]


def _valid_month_header(col: str) -> bool:
    try:
        ym = str(col)[:7]
        pd.to_datetime(ym + "-01")
        return True
    except Exception:
        return False


def _month_date_from_col(col: str) -> pd.Timestamp:
    return pd.to_datetime(str(col)[:7] + "-01")


def _month_headers_costs_valid(df: pd.DataFrame) -> list[str]:
    cols = []
    for c in df.columns:
        s = str(c)
        if len(s) >= 7 and s.endswith((" Costs", "Costs")) and _valid_month_header(s):
            cols.append(s)
    return sorted(set(cols), key=_month_date_from_col)


def _month_headers_all_valid(df: pd.DataFrame) -> list[str]:
    cols = []
    for c in df.columns:
        s = str(c)
        if len(s) >= 7 and s.endswith(
            (" Costs", "Costs", " Forecast", "Forecast", " Forecast Snapshot Costs")
        ):
            if _valid_month_header(s):
                cols.append(s)
    return sorted(set(cols), key=_month_date_from_col)


def path_check(file_path: Path) -> pd.DataFrame:
    logging.info("Starting PathCheck for %s", Path(file_path).name)
    a = _read_sheet(file_path, SHEET_ACTUALS_ETC)

    if "Path ID" not in a.columns:
        raise ValueError("Column 'Path ID' not found in 'Snapshot Actuals plus ETC'.")

    for k in ["Project ID", "Project Name"]:
        if k not in a.columns:
            raise ValueError(f"Missing '{k}' in 'Snapshot Actuals plus ETC'.")

    out = a[["Project ID", "Project Name", "Path ID"]].copy()

    out["Project ID"] = out["Project ID"].astype(str).str.strip()
    out["Project Name"] = out["Project Name"].astype(str).str.strip()

    out = out[
        (out["Project ID"] != "Total")
        & out["Path ID"].notna()
        & (out["Path ID"].astype(str).str.strip() != "")
    ].copy()

    out["Path ID"] = out["Path ID"].astype(str).str.strip()

    out["Path_Segments"] = out["Path ID"].str.split(".").apply(len)
    out["Expected_Segments"] = 8
    out["Path_Variance"] = out["Path_Segments"] - out["Expected_Segments"]
    out["Path_Flag"] = out["Path_Segments"].apply(lambda x: "OK" if x == 8 else "Mismatch")

    out = out.sort_values(["Project ID", "Project Name"]).reset_index(drop=True)
    logging.info("Completed PathCheck for %s", Path(file_path).name)
    return out


def snapshot_check(file_path: Path, actuals_month: str, forecast_start: str) -> pd.DataFrame:
    logging.info(
        "Starting SnapshotCheck for %s | actuals_month=%s | forecast_start=%s",
        Path(file_path).name, actuals_month, forecast_start
    )

    act_cut = _to_month(actuals_month)
    fc_start = _to_month(forecast_start)

    ob = _read_sheet(file_path, SHEET_OB)
    cb = _read_sheet(file_path, SHEET_CB)
    ch = _read_sheet(file_path, SHEET_CHANGE)
    a = _read_sheet(file_path, SHEET_ACTUALS_ETC)

    ob_cols = [
        c for c in _month_cols(ob, " Costs")
        if pd.to_datetime(RE_YM.match(c)["ym"] + "-01") >= pd.Timestamp("2018-01-01")
    ]
    cb_cols = [
        c for c in _month_cols(cb, " Costs")
        if pd.to_datetime(RE_YM.match(c)["ym"] + "-01") >= pd.Timestamp("2018-01-01")
    ]
    ch_cols = [
        c for c in _month_cols(ch, " Costs")
        if pd.to_datetime(RE_YM.match(c)["ym"] + "-01") >= pd.Timestamp("2018-01-01")
    ]

    ob_g = _sum_month_cols_by_project(ob, ob_cols, "OB_Total")
    cb_g = _sum_month_cols_by_project(cb, cb_cols, "CB_Total")
    ch_g = _sum_month_cols_by_project(ch, ch_cols, "Changes")

    for k in ["Project ID", "Project Name"]:
        if k not in a.columns:
            raise ValueError(f"Missing '{k}' in 'Snapshot Actuals plus ETC'.")

    a["Project ID"] = a["Project ID"].astype(str).str.strip()
    a = a[a["Project ID"] != "Total"].copy()

    act_cols_all = _month_cols(a, " Actuals Costs")
    fc_cols_all = _month_cols(a, " Forecast Snapshot Costs")

    if not act_cols_all:
        raise ValueError("No 'YYYY-MM Actuals Costs' columns found in 'Snapshot Actuals plus ETC'.")
    if not fc_cols_all:
        raise ValueError("No 'YYYY-MM Forecast Snapshot Costs' columns found in 'Snapshot Actuals plus ETC'.")

    act_keep = [c for c in act_cols_all if pd.to_datetime(RE_YM.match(c)["ym"] + "-01") <= act_cut]
    for c in act_keep:
        a[c] = pd.to_numeric(a[c], errors="coerce").fillna(0.0)
    a["Actuals"] = a[act_keep].sum(axis=1) if act_keep else pd.Series(0.0, index=a.index)

    monthly_col = f"{actuals_month} Actuals Costs"
    if monthly_col in a.columns:
        a["MonthlyActuals"] = pd.to_numeric(a[monthly_col], errors="coerce").fillna(0.0)
    else:
        a["MonthlyActuals"] = 0.0

    fc_keep = [c for c in fc_cols_all if pd.to_datetime(RE_YM.match(c)["ym"] + "-01") >= fc_start]
    for c in fc_keep:
        a[c] = pd.to_numeric(a[c], errors="coerce").fillna(0.0)
    a["Forecast"] = a[fc_keep].sum(axis=1) if fc_keep else pd.Series(0.0, index=a.index)

    a_g = a.groupby(["Project ID", "Project Name"], as_index=False)[["Actuals", "MonthlyActuals", "Forecast"]].sum()

    def _keys(df: pd.DataFrame) -> pd.DataFrame:
        return df[["Project ID"]].drop_duplicates()

    master = pd.concat([_keys(ob_g), _keys(cb_g), _keys(a_g), _keys(ch_g)], ignore_index=True).drop_duplicates()

    master = master.merge(a_g[["Project ID", "Project Name"]], on="Project ID", how="left")
    master = master.merge(
        ob_g[["Project ID", "Project Name"]].rename(columns={"Project Name": "Name_OB"}),
        on="Project ID",
        how="left",
    )
    master = master.merge(
        cb_g[["Project ID", "Project Name"]].rename(columns={"Project Name": "Name_CB"}),
        on="Project ID",
        how="left",
    )
    master["Project Name"] = master["Project Name"].fillna(master["Name_OB"]).fillna(master["Name_CB"])
    master = master.drop(columns=["Name_OB", "Name_CB"])

    out = master.merge(ob_g[["Project ID", "OB_Total"]], on="Project ID", how="left")
    out = out.merge(cb_g[["Project ID", "CB_Total"]], on="Project ID", how="left")
    out = out.merge(a_g[["Project ID", "Actuals", "MonthlyActuals", "Forecast"]], on="Project ID", how="left")
    out = out.merge(ch_g[["Project ID", "Changes"]], on="Project ID", how="left")

    for c in ["OB_Total", "CB_Total", "Actuals", "MonthlyActuals", "Forecast", "Changes"]:
        out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0.0)

    out["Changes_Check"] = out["OB_Total"] + out["Changes"]
    out["Changes_Variance"] = out["CB_Total"] - out["Changes_Check"]
    out["Changes_Flag"] = out["Changes_Variance"].abs().apply(lambda x: "OK" if x < 1 else "Mismatch")
    out["EAC"] = out["Actuals"] + out["Forecast"]

    out = out[
        [
            "Project ID",
            "Project Name",
            "OB_Total",
            "CB_Total",
            "Actuals",
            "MonthlyActuals",
            "Forecast",
            "Changes",
            "Changes_Check",
            "Changes_Variance",
            "Changes_Flag",
            "EAC",
        ]
    ].sort_values(["Project ID", "Project Name"]).reset_index(drop=True)

    logging.info("Completed SnapshotCheck for %s", Path(file_path).name)
    return out


def _actuals_by_month(file_path: Path) -> pd.DataFrame:
    logging.info("Building Actuals by month for %s", Path(file_path).name)
    df = _read_sheet(file_path, SHEET_ACTUALS)
    month_cols = _month_headers_costs_valid(df)

    if not month_cols:
        return pd.DataFrame(columns=["Month", "Actuals"])

    tmp = df[month_cols].copy()
    for c in month_cols:
        tmp[c] = pd.to_numeric(tmp[c], errors="coerce").fillna(0.0)

    unp = tmp.melt(var_name="MonthRaw", value_name="Amount")
    unp["Month"] = pd.to_datetime(unp["MonthRaw"].str[:7] + "-01", errors="coerce")

    out = (
        unp.groupby("Month", as_index=False)["Amount"]
        .sum()
        .rename(columns={"Amount": "Actuals"})
        .sort_values("Month")
        .reset_index(drop=True)
    )
    return out


def _eac_total(file_path: Path) -> pd.DataFrame:
    logging.info("Building EAC total for %s", Path(file_path).name)
    df = _read_sheet(file_path, SHEET_ACTUALS_ETC)
    month_cols = _month_headers_all_valid(df)

    if not month_cols:
        return pd.DataFrame({"EAC": [0.0]})

    total = 0.0
    for c in month_cols:
        total += pd.to_numeric(df[c], errors="coerce").fillna(0.0).sum()

    return pd.DataFrame({"EAC": [round(float(total), 2)]})


def _eac_by_month(file_path: Path) -> pd.DataFrame:
    logging.info("Building EAC by month for %s", Path(file_path).name)
    df = _read_sheet(file_path, SHEET_ACTUALS_ETC)
    month_cols = _month_headers_all_valid(df)

    if not month_cols:
        return pd.DataFrame(columns=["Month", "EAC_Month"])

    tmp = df[month_cols].copy()
    for c in month_cols:
        tmp[c] = pd.to_numeric(tmp[c], errors="coerce").fillna(0.0)

    unp = tmp.melt(var_name="MonthRaw", value_name="Amount")
    unp["Month"] = pd.to_datetime(unp["MonthRaw"].str[:7] + "-01", errors="coerce")

    out = (
        unp.groupby("Month", as_index=False)["Amount"]
        .sum()
        .rename(columns={"Amount": "EAC_Month"})
        .sort_values("Month")
        .reset_index(drop=True)
    )
    return out


def snapshotcheck_compare(snap_a: pd.DataFrame, snap_b: pd.DataFrame) -> pd.DataFrame:
    logging.info("Building SnapshotCheck comparison")
    keep = ["Project ID", "Project Name", "OB_Total", "CB_Total", "Actuals", "Forecast", "EAC"]

    a0 = snap_a[[c for c in keep if c in snap_a.columns]].copy()
    b0 = snap_b[[c for c in keep if c in snap_b.columns]].copy()

    merged = pd.merge(
        a0,
        b0,
        on=["Project ID"],
        how="outer",
        suffixes=("", "_B"),
    )

    merged["Project Name"] = merged["Project Name"].combine_first(merged.get("Project Name_B"))

    if "Project Name_B" in merged.columns:
        merged = merged.drop(columns=["Project Name_B"])

    numeric_cols = [
        "OB_Total", "CB_Total", "Actuals", "Forecast", "EAC",
        "OB_Total_B", "CB_Total_B", "Actuals_B", "Forecast_B", "EAC_B",
    ]
    for c in numeric_cols:
        if c not in merged.columns:
            merged[c] = 0.0
        merged[c] = pd.to_numeric(merged[c], errors="coerce").fillna(0.0)

    merged["OB_Delta"] = (merged["OB_Total"] - merged["OB_Total_B"]).round(2)
    merged["CB_Delta"] = (merged["CB_Total"] - merged["CB_Total_B"]).round(2)
    merged["Actuals_Delta"] = (merged["Actuals"] - merged["Actuals_B"]).round(2)
    merged["Forecast_Delta"] = (merged["Forecast"] - merged["Forecast_B"]).round(2)
    merged["EAC_Delta"] = (merged["EAC"] - merged["EAC_B"]).round(2)

    out = merged[
        ["Project ID", "Project Name", "OB_Delta", "CB_Delta", "Actuals_Delta", "Forecast_Delta", "EAC_Delta"]
    ].sort_values("Project ID").reset_index(drop=True)

    return out


def actuals_compare_by_month(file_a: Path, file_b: Path) -> pd.DataFrame:
    logging.info("Building Actuals compare by month")
    a = _actuals_by_month(file_a).rename(columns={"Actuals": "Actuals_A"})
    b = _actuals_by_month(file_b).rename(columns={"Actuals": "Actuals_B"})

    merged = pd.merge(a, b, on="Month", how="outer")
    merged["Actuals_A"] = pd.to_numeric(merged["Actuals_A"], errors="coerce").fillna(0.0)
    merged["Actuals_B"] = pd.to_numeric(merged["Actuals_B"], errors="coerce").fillna(0.0)
    merged["Actuals_Delta"] = (merged["Actuals_A"] - merged["Actuals_B"]).round(2)

    merged = merged.sort_values("Month").reset_index(drop=True)
    merged["Month"] = merged["Month"].dt.strftime("%m/%Y")
    return merged


def eac_compare_total(file_a: Path, file_b: Path) -> pd.DataFrame:
    logging.info("Building EAC total comparison")
    a = _eac_total(file_a)
    b = _eac_total(file_b)

    eac_a = float(a.loc[0, "EAC"]) if not a.empty else 0.0
    eac_b = float(b.loc[0, "EAC"]) if not b.empty else 0.0

    return pd.DataFrame({
        "EAC_A": [round(eac_a, 2)],
        "EAC_B": [round(eac_b, 2)],
        "EAC_Delta": [round(eac_a - eac_b, 2)],
    })


def eac_compare_by_month(file_a: Path, file_b: Path) -> pd.DataFrame:
    logging.info("Building EAC compare by month")
    a = _eac_by_month(file_a).rename(columns={"EAC_Month": "EAC_A"})
    b = _eac_by_month(file_b).rename(columns={"EAC_Month": "EAC_B"})

    merged = pd.merge(a, b, on="Month", how="outer")
    merged["EAC_A"] = pd.to_numeric(merged["EAC_A"], errors="coerce").fillna(0.0)
    merged["EAC_B"] = pd.to_numeric(merged["EAC_B"], errors="coerce").fillna(0.0)
    merged["EAC_Delta"] = (merged["EAC_A"] - merged["EAC_B"]).round(2)

    merged = merged.sort_values("Month").reset_index(drop=True)
    merged["Month"] = merged["Month"].dt.strftime("%m/%Y")
    return merged


def write_outputs(out_path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    logging.info("Writing outputs to %s", Path(out_path).name)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            safe_name = str(sheet_name)[:31]
            if df is None:
                pd.DataFrame().to_excel(writer, sheet_name=safe_name, index=False)
            else:
                df.to_excel(writer, sheet_name=safe_name, index=False)

    logging.info("Finished writing outputs to %s", Path(out_path).name)
